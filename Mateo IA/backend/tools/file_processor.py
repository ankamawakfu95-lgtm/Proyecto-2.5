"""Carga y generación de archivos (PDF, Word, Excel, texto/markdown).

Sigue la misma filosofía que el resto del proyecto: cada dependencia es
opcional. Si una librería no está instalada, la función correspondiente
devuelve un mensaje claro en vez de romper el proceso.
"""
from __future__ import annotations

import csv
import io
from datetime import datetime
from pathlib import Path
from typing import Any, Dict

MAX_FILE_SIZE_MB = 25
MAX_EXTRACTED_CHARS = 40_000  # evita saturar el contexto del LLM con un PDF de 300 páginas

GENERATED_DIR = Path(__file__).resolve().parents[1] / "data" / "generated"
GENERATED_DIR.mkdir(parents=True, exist_ok=True)


class FileProcessingError(Exception):
    """Error controlado al leer o generar un archivo."""


# ==========================================
# CARGA / EXTRACCIÓN DE TEXTO
# ==========================================
def load_file(path: str | Path, original_name: str | None = None) -> Dict[str, Any]:
    """Extrae texto de un archivo soportado. Devuelve {text, kind, truncated}."""
    path = Path(path)
    name = original_name or path.name
    size_mb = path.stat().st_size / (1024 * 1024)
    if size_mb > MAX_FILE_SIZE_MB:
        raise FileProcessingError(f"El archivo pesa {size_mb:.1f}MB; el límite es {MAX_FILE_SIZE_MB}MB.")

    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            text = _extract_pdf(path)
        elif suffix == ".docx":
            text = _extract_docx(path)
        elif suffix in (".xlsx", ".xls"):
            text = _extract_xlsx(path)
        elif suffix == ".csv":
            text = _extract_csv(path)
        elif suffix in (".txt", ".md", ".py", ".js", ".ts", ".json", ".log"):
            text = path.read_text(encoding="utf-8", errors="replace")
        else:
            raise FileProcessingError(f"Tipo de archivo no soportado: {suffix or '(sin extensión)'}")
    except FileProcessingError:
        raise
    except Exception as e:  # dependencia faltante, archivo corrupto, etc.
        raise FileProcessingError(f"No pude leer '{name}': {e}") from e

    truncated = len(text) > MAX_EXTRACTED_CHARS
    if truncated:
        text = text[:MAX_EXTRACTED_CHARS]

    return {"text": text.strip(), "kind": suffix.lstrip("."), "truncated": truncated, "name": name}


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise FileProcessingError("Falta 'pypdf'. Instalá con: pip install pypdf") from e
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as e:
        raise FileProcessingError("Falta 'python-docx'. Instalá con: pip install python-docx") from e
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError as e:
        raise FileProcessingError("Falta 'openpyxl'. Instalá con: pip install openpyxl") from e
    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"### Hoja: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                lines.append(" | ".join("" if c is None else str(c) for c in row))
    return "\n".join(lines)


def _extract_csv(path: Path) -> str:
    with path.open("r", encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.reader(f)
        return "\n".join(" | ".join(row) for row in reader)


# ==========================================
# GENERACIÓN DE ARCHIVOS
# ==========================================
def generate_document(kind: str, title: str, content: str) -> Path:
    """Genera un archivo (docx/pdf/xlsx/md/txt) a partir de texto y lo guarda en data/generated/."""
    kind = kind.lower().lstrip(".")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    # Sin espacios ni caracteres especiales: el nombre viaja en una URL (/files/{filename}).
    safe_title = "".join(c if c.isalnum() or c in "-_" else "_" for c in title.replace(" ", "_")).strip("_")[:60] or "documento"
    filename = f"{safe_title}_{timestamp}.{kind}"
    target = GENERATED_DIR / filename

    if kind == "docx":
        _write_docx(target, title, content)
    elif kind == "pdf":
        _write_pdf(target, title, content)
    elif kind in ("xlsx", "excel"):
        target = target.with_suffix(".xlsx")
        _write_xlsx(target, title, content)
    elif kind in ("md", "markdown"):
        target.write_text(f"# {title}\n\n{content}\n", encoding="utf-8")
    elif kind == "txt":
        target.write_text(f"{title}\n\n{content}\n", encoding="utf-8")
    else:
        raise FileProcessingError(f"No sé generar archivos '.{kind}'. Formatos soportados: docx, pdf, xlsx, md, txt.")

    return target


def _write_docx(target: Path, title: str, content: str) -> None:
    try:
        import docx
    except ImportError as e:
        raise FileProcessingError("Falta 'python-docx'. Instalá con: pip install python-docx") from e
    document = docx.Document()
    document.add_heading(title, level=1)
    for paragraph in content.split("\n\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())
    document.save(str(target))


def _write_pdf(target: Path, title: str, content: str) -> None:
    try:
        from fpdf import FPDF
    except ImportError as e:
        raise FileProcessingError("Falta 'fpdf2'. Instalá con: pip install fpdf2") from e
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, title)
    pdf.ln(4)
    pdf.set_font("Helvetica", "", 11)
    for paragraph in content.split("\n\n"):
        if paragraph.strip():
            pdf.multi_cell(0, 7, paragraph.strip())
            pdf.ln(2)
    pdf.output(str(target))


def _write_xlsx(target: Path, title: str, content: str) -> None:
    try:
        import openpyxl
    except ImportError as e:
        raise FileProcessingError("Falta 'openpyxl'. Instalá con: pip install openpyxl") from e
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Datos"
    # Si el contenido parece tabular (filas separadas por "|"), lo respeta; si no, una fila por párrafo.
    for line in content.splitlines():
        if "|" in line:
            sheet.append([cell.strip() for cell in line.split("|")])
        elif line.strip():
            sheet.append([line.strip()])
    workbook.save(str(target))
